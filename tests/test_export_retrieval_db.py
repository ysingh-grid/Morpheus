"""End-to-end coverage for exporting the populated retrieval database to Excel."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scripts.export_retrieval_db import EXPORT_TABLES, export_retrieval_database


def test_export_retrieval_database_writes_all_normalized_tables(tmp_path: Path) -> None:
    """Export the configured database and preserve each normalized table as a worksheet."""
    output_path = tmp_path / "retrieval-export.xlsx"

    row_counts = export_retrieval_database(output_path)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    assert workbook.sheetnames == list(EXPORT_TABLES)
    for table_name in EXPORT_TABLES:
        worksheet = workbook[table_name]
        assert worksheet.max_row == row_counts[table_name] + 1
    assert row_counts["documents"] >= 1
    assert row_counts["children"] >= 1
    assert row_counts["parents"] >= 1
