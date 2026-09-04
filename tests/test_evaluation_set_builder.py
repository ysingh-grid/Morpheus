"""Live database coverage for building the human-review retrieval benchmark."""

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from retrieval.build_evaluation_set import build_review_cases, import_human_review, write_review_set


def test_build_review_cases_creates_one_hundred_grounded_review_rows(tmp_path: Path) -> None:
    """Build the IFC review matrix and retain its categories in JSON and Excel outputs."""
    cases = build_review_cases()
    counts = Counter(case["category"] for case in cases)
    json_output = tmp_path / "review.json"
    workbook_output = tmp_path / "review.xlsx"

    write_review_set(cases, json_output, workbook_output)

    assert len(cases) == 100
    assert counts == {"text": 30, "table": 30, "figure": 20, "numeric": 10, "ambiguous": 5, "no_answer": 5}
    assert all(case["verification_status"] == "pending_human_review" for case in cases)
    assert json_output.exists()
    workbook = load_workbook(workbook_output, read_only=True)
    assert workbook["review_cases"].max_row == 101

    editable_workbook = load_workbook(workbook_output)
    worksheet = editable_workbook["review_cases"]
    worksheet.cell(row=2, column=10, value="human_verified")
    worksheet.cell(row=2, column=11, value="Reviewer")
    editable_workbook.save(workbook_output)
    assert import_human_review(workbook_output, json_output) == 100
