"""Build an annotation-ready, 100-case retrieval evaluation set from PostgreSQL."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from pathlib import Path
from typing import Any, TypedDict

import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from retrieval.pg_engine import _database_url

DEFAULT_DOCUMENT_PATH = Path("dummy_data/ifc-annual-report-2024-financials.pdf")
DEFAULT_JSON_OUTPUT = Path("retrieval/evaluation_cases_ifc_2024_human_review.json")
DEFAULT_WORKBOOK_OUTPUT = Path("dummy_data/ifc-annual-report-2024-financials.evaluation-review.xlsx")
REVIEW_HEADERS = (
    "id",
    "category",
    "query",
    "expected_status",
    "expected_parent_id",
    "expected_table_id",
    "expected_figure_id",
    "expected_answer",
    "source_excerpt",
    "verification_status",
    "reviewer",
    "reviewer_notes",
)


class ReviewCase(TypedDict):
    """One source-grounded evaluation case awaiting human verification."""

    id: str
    category: str
    query: str
    expected_status: str
    expected_parent_id: str
    expected_table_id: str
    expected_figure_id: str
    expected_answer: str
    source_excerpt: str
    verification_status: str
    reviewer: str
    reviewer_notes: str


def _section_title(parent_text: str) -> str:
    """Extract a concise section title from a normalized parent block."""
    for line in parent_text.splitlines():
        candidate = line.lstrip("#").strip()
        if candidate and not candidate.startswith("["):
            return candidate[:120]
    return parent_text.replace("\n", " ")[:120]


def _numeric_prompt(markdown: str, heading: str) -> tuple[str, str]:
    """Extract one row label and numeric value to make a table-value question."""
    for line in markdown.splitlines():
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 2 or all(set(cell) <= {"-", ":", " "} for cell in cells):
            continue
        value_cells = [cell for cell in cells[1:] if re.search(r"\d", cell)]
        if value_cells and cells[0] and not re.search(r"^20\d\d$", cells[0]):
            return (
                f"What value is reported for {cells[0][:100]} in {heading[:100]}?",
                value_cells[0][:160],
            )
    return (f"What numeric values are reported in {heading[:100]}?", "Human reviewer must identify the answer.")


def _fetch_rows(document_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Read parent, table, and figure records belonging to one uploaded document."""
    with psycopg.connect(_database_url()) as connection:
        with connection.cursor(row_factory=psycopg.rows.dict_row) as cursor:
            cursor.execute(
                "SELECT id, parent_text FROM parents WHERE doc_id = %s ORDER BY id", (document_id,)
            )
            parents = list(cursor.fetchall())
            cursor.execute(
                """
                SELECT source_table.id, source_table.heading, source_table.markdown, parent.id AS parent_id
                FROM tables AS source_table
                JOIN parents AS parent
                  ON parent.doc_id = source_table.doc_id
                 AND source_table.id = ANY(parent.table_ids)
                WHERE source_table.doc_id = %s
                ORDER BY source_table.id
                """,
                (document_id,),
            )
            tables = list(cursor.fetchall())
            cursor.execute(
                """
                SELECT figure.id, figure.caption, parent.id AS parent_id
                FROM figures AS figure
                JOIN parents AS parent
                  ON parent.doc_id = figure.doc_id
                 AND figure.id = ANY(parent.figure_ids)
                WHERE figure.doc_id = %s
                ORDER BY figure.id
                """,
                (document_id,),
            )
            figures = list(cursor.fetchall())
    return parents, tables, figures


def _default_document_id() -> str:
    """Resolve the evaluation corpus through the same content-addressed ID contract."""
    configured_id = os.getenv("EVALUATION_DOCUMENT_ID")
    if configured_id:
        return configured_id
    digest = hashlib.sha256()
    with DEFAULT_DOCUMENT_PATH.open("rb") as source_file:
        for block in iter(lambda: source_file.read(1_048_576), b""):
            digest.update(block)
    return f"sha256-{digest.hexdigest()}"


def build_review_cases(document_id: str | None = None) -> list[ReviewCase]:
    """Generate exactly 100 source-linked cases, each pending human verification."""
    document_id = document_id or _default_document_id()
    parents, tables, figures = _fetch_rows(document_id)
    if len(parents) < 30 or len(tables) < 30 or len(figures) < 20:
        raise RuntimeError("The document lacks enough parent, table, or figure records for the 100-case set.")

    cases: list[ReviewCase] = []

    def add_case(**values: str) -> None:
        cases.append(
            {
                "verification_status": "pending_human_review",
                "reviewer": "",
                "reviewer_notes": "",
                **values,
            }
        )

    for index, parent in enumerate(parents[:30], start=1):
        title = _section_title(parent["parent_text"])
        add_case(
            id=f"text-{index:03d}",
            category="text",
            query=f"What does the IFC report state about {title}?",
            expected_status="grounded",
            expected_parent_id=parent["id"],
            expected_table_id="",
            expected_figure_id="",
            expected_answer="Human reviewer must verify the supporting passage and answer.",
            source_excerpt=title,
        )

    for index, table in enumerate(tables[:30], start=1):
        heading = table["heading"] or table["id"]
        add_case(
            id=f"table-{index:03d}",
            category="table",
            query=f"What information is reported in {heading}?",
            expected_status="grounded",
            expected_parent_id=table["parent_id"],
            expected_table_id=table["id"],
            expected_figure_id="",
            expected_answer="Human reviewer must verify the supporting rows and answer.",
            source_excerpt=heading[:240],
        )

    for index, figure in enumerate(figures[:20], start=1):
        caption = figure["caption"].replace("\n", " ")
        add_case(
            id=f"figure-{index:03d}",
            category="figure",
            query=f"Which figure shows: {caption[:180]}?",
            expected_status="grounded",
            expected_parent_id=figure["parent_id"],
            expected_table_id="",
            expected_figure_id=figure["id"],
            expected_answer="Human reviewer must verify that the caption matches the source figure.",
            source_excerpt=caption[:500],
        )

    for index, table in enumerate(tables[30:40], start=1):
        heading = table["heading"] or table["id"]
        query, expected_answer = _numeric_prompt(table["markdown"], heading)
        add_case(
            id=f"numeric-{index:03d}",
            category="numeric",
            query=query,
            expected_status="grounded",
            expected_parent_id=table["parent_id"],
            expected_table_id=table["id"],
            expected_figure_id="",
            expected_answer=expected_answer,
            source_excerpt=heading[:240],
        )

    for index, query in enumerate(
        ("What was the net income?", "What was the portfolio value?", "What were the reserves?", "What were total assets?", "What were disbursements?"),
        start=1,
    ):
        add_case(
            id=f"ambiguous-{index:03d}",
            category="ambiguous",
            query=query,
            expected_status="clarification_needed",
            expected_parent_id="",
            expected_table_id="",
            expected_figure_id="",
            expected_answer="Reviewer must confirm whether clarification is required.",
            source_excerpt="Potentially multiple reporting periods, tables, or definitions apply.",
        )

    for index, query in enumerate(
        (
            "What was IFC's Mars colonization budget?",
            "Which IFC table reports the price of Bitcoin?",
            "What was IFC's 2035 lunar lending target?",
            "Which chart shows IFC's social media followers?",
            "What was the annual rainfall at IFC headquarters?",
        ),
        start=1,
    ):
        add_case(
            id=f"no-answer-{index:03d}",
            category="no_answer",
            query=query,
            expected_status="not_found",
            expected_parent_id="",
            expected_table_id="",
            expected_figure_id="",
            expected_answer="No answer should be returned from this uploaded document.",
            source_excerpt="Out-of-scope question.",
        )

    if len(cases) != 100:
        raise AssertionError(f"Expected 100 evaluation cases, generated {len(cases)}.")
    return cases


def write_review_set(cases: list[ReviewCase], json_output: Path, workbook_output: Path) -> None:
    """Write the machine-readable matrix and a human-reviewable Excel workbook."""
    payload = {
        "name": "IFC 2024 financials human-review retrieval benchmark",
        "top_k": 5,
        "require_human_verification": True,
        "minimum_human_verified_cases": 100,
        "cases": cases,
    }
    json_output.parent.mkdir(parents=True, exist_ok=True)
    json_output.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    workbook_output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "review_cases"
    worksheet.append(REVIEW_HEADERS)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    worksheet.freeze_panes = "A2"
    for case in cases:
        worksheet.append([case[header] for header in REVIEW_HEADERS])
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(workbook_output)


def import_human_review(workbook_path: Path, json_output: Path) -> int:
    """Merge reviewer edits from the workbook into the matching JSON case matrix."""
    payload = json.loads(json_output.read_text(encoding="utf-8"))
    cases_by_id = {case["id"]: case for case in payload["cases"]}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["review_cases"]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    if tuple(headers) != REVIEW_HEADERS:
        raise ValueError("Review workbook headers do not match the expected evaluation-review format.")

    updated = 0
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(REVIEW_HEADERS, values, strict=True))
        case = cases_by_id.get(row["id"])
        if case is None:
            raise ValueError(f"Review workbook contains unknown case ID: {row['id']}")
        for header in REVIEW_HEADERS[3:]:
            case[header] = row[header] or ""
        updated += 1
    json_output.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return updated


def main() -> None:
    """Build a 100-case review set for one loaded document."""
    parser = argparse.ArgumentParser(description="Build an annotation-ready 100-case retrieval benchmark.")
    parser.add_argument(
        "--document-id",
        default=None,
        help="PostgreSQL document ID to sample; defaults to the IFC file's content ID.",
    )
    parser.add_argument("--json-output", type=Path, default=DEFAULT_JSON_OUTPUT, help="JSON matrix output.")
    parser.add_argument("--workbook-output", type=Path, default=DEFAULT_WORKBOOK_OUTPUT, help="Excel review workbook output.")
    parser.add_argument(
        "--import-review-workbook",
        type=Path,
        help="Merge completed reviewer edits from this workbook into --json-output.",
    )
    arguments = parser.parse_args()
    if arguments.import_review_workbook:
        updated = import_human_review(arguments.import_review_workbook, arguments.json_output)
        print(f"Imported review edits for {updated} cases into {arguments.json_output}.")
        return
    cases = build_review_cases(arguments.document_id)
    write_review_set(cases, arguments.json_output, arguments.workbook_output)
    print(f"Created {len(cases)} pending-human-review cases.")
    print(f"JSON: {arguments.json_output}")
    print(f"Workbook: {arguments.workbook_output}")


if __name__ == "__main__":
    main()
